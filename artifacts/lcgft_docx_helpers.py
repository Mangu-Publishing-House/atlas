"""Shared helpers for the MANGU LCGFT master doc builder."""

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = RGBColor(0x1B, 0x36, 0x5D)
NAVY2 = RGBColor(0x14, 0x28, 0x48)
GOLD = RGBColor(0xB8, 0x94, 0x4A)
INK = RGBColor(0x1C, 0x1C, 0x1E)
SLATE = RGBColor(0x4E, 0x5A, 0x68)
MUTED = RGBColor(0x6B, 0x73, 0x80)
PALE = "F3F5F8"


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=INK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def _tc_insert(tcPr, tag, element):
    """Insert element into tcPr in a schema-safe way: before later siblings if present."""
    order = [
        "cnfStyle","tcW","gridSpan","hMerge","vMerge","tcBorders","shd",
        "noWrap","tcMar","textDirection","tcFitText","vAlign","hideMark","headers",
    ]
    el_tag = tag
    # remove existing same tag
    for child in list(tcPr):
        if child.tag == qn(f"w:{el_tag}"):
            tcPr.remove(child)
    idx_target = order.index(el_tag) if el_tag in order else 99
    insert_at = None
    for i, child in enumerate(list(tcPr)):
        local = child.tag.split("}")[-1]
        if local in order and order.index(local) > idx_target:
            insert_at = i
            break
    if insert_at is None:
        tcPr.append(element)
    else:
        tcPr.insert(insert_at, element)


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    _tc_insert(tcPr, "shd", shd)


def set_cell_borders(cell, color="C5CAD3", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    _tc_insert(tcPr, "tcBorders", tcBorders)


def set_cell_margins(cell, top=40, bottom=40, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    _tc_insert(tcPr, "tcMar", tcMar)


def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))


def set_table_widths(table, widths):
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else tbl.get_or_add_tblPr()
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for child in list(grid):
            grid.remove(child)
    else:
        grid = OxmlElement("w:tblGrid")
        tblPr.addnext(grid)
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(w))
        grid.append(gc)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths):
                cell.width = widths[i]
                tcPr = cell._tc.get_or_add_tcPr()
                tcW = tcPr.find(qn("w:tcW"))
                if tcW is None:
                    tcW = OxmlElement("w:tcW")
                    tcPr.append(tcW)
                tcW.set(qn("w:w"), str(widths[i]))
                tcW.set(qn("w:type"), "dxa")


def add_bottom_border_para(paragraph, color="B8944A", sz="12"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    order = [
        "pStyle","keepNext","keepLines","pageBreakBefore","framePr","widowControl",
        "numPr","suppressLineNumbers","pBdr","shd","tabs","suppressAutoHyphens",
        "kinsoku","wordWrap","overflowPunct","topLinePunct","autoSpaceDE",
        "autoSpaceDN","bidi","adjustRightInd","snapToGrid","spacing","ind",
        "contextualSpacing","mirrorIndents","suppressOverlap","jc","textDirection",
        "textAlignment","textboxTightWrap","outlineLvl","divId","cnfStyle",
        "rPr","sectPr","pPrChange",
    ]
    for child in list(pPr):
        if child.tag == qn("w:pBdr"):
            pPr.remove(child)
    idx_target = order.index("pBdr")
    insert_at = None
    for i, child in enumerate(list(pPr)):
        local = child.tag.split("}")[-1]
        if local in order and order.index(local) > idx_target:
            insert_at = i
            break
    if insert_at is None:
        pPr.append(pBdr)
    else:
        pPr.insert(insert_at, pBdr)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instr)
    run._r.append(fldChar2)


def add_toc_field(paragraph):
    run = paragraph.add_run()
    r = run._r
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r' TOC \o "1-2" \h \z \u '
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Right-click and choose Update Field to refresh this table of contents."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r.append(begin)
    r.append(instr)
    r.append(sep)
    r.append(text)
    r.append(end)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.15
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    rFonts.set(qn("w:eastAsia"), "Calibri")
    for i, (size, before, after, color) in enumerate(
        [(22, 0, 10, NAVY), (16, 16, 8, NAVY), (13, 12, 6, NAVY2)], start=1
    ):
        s = doc.styles[f"Heading {i}"]
        s.font.name = "Calibri"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = color
        s.font.italic = False
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True


def add_header_footer(section):
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.clear()
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = hp.add_run("MANGU  ·  LCGFT Master Taxonomy")
    set_run_font(r, "Calibri", 9, True, False, NAVY)
    r2 = hp.add_run("    |    Internal catalog authority  ·  v1.0  ·  15 September 2026")
    set_run_font(r2, "Calibri", 9, False, False, MUTED)
    add_bottom_border_para(hp, "1B365D", "12")
    hp.paragraph_format.space_after = Pt(6)
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.clear()
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = fp.add_run(
        "Source: Library of Congress Genre/Form Terms  ·  official list updated December 2025  ·  compiled 19 August 2026"
    )
    set_run_font(r, "Calibri", 8, False, False, MUTED)
    r = fp.add_run("          p. ")
    set_run_font(r, "Calibri", 8, False, False, MUTED)
    add_page_number(fp)
    for run in fp.runs:
        if run.font.size is None:
            set_run_font(run, "Calibri", 8, False, False, MUTED)


def p(doc, text, size=11, bold=False, italic=False, color=INK, space_after=8, space_before=0, align="left"):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.line_spacing = 1.15
    if align == "center":
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = para.add_run(text)
    set_run_font(run, "Calibri", size, bold, italic, color)
    return para


def bullet(doc, text):
    para = doc.add_paragraph(style="List Bullet")
    para.clear()
    para.paragraph_format.left_indent = Inches(0.3)
    para.paragraph_format.space_after = Pt(3)
    para.paragraph_format.space_before = Pt(0)
    run = para.add_run(text)
    set_run_font(run, "Calibri", 11, False, False, INK)
    return para


def tree_line(doc, text, level, extra=""):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(1)
    para.paragraph_format.line_spacing = 1.05
    para.paragraph_format.left_indent = Inches(0.16 * level)
    if level == 0:
        size, bold, color = 11, True, NAVY
        prefix = ""
    else:
        size, bold, color = 10, False, INK
        prefix = ("│  " * (level - 1)) + "├─ "
    run = para.add_run(prefix + text)
    set_run_font(run, "Calibri", size, bold, False, color)
    if extra:
        run2 = para.add_run(extra)
        set_run_font(run2, "Calibri", 9, False, True, MUTED)
    return para


def simple_table(doc, headers, rows, col_widths, header_fill="1B365D"):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, col_widths)
    hdr = table.rows[0]
    set_repeat_header(hdr)
    prevent_row_split(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ""
        shade_cell(cell, header_fill)
        set_cell_margins(cell, 50, 50, 70, 70)
        set_cell_borders(cell, "1B365D", "4")
        para = cell.paragraphs[0]
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.space_before = Pt(0)
        run = para.add_run(h)
        set_run_font(run, "Calibri", 9, True, False, RGBColor(255, 255, 255))
    for r_i, row in enumerate(rows):
        rr = table.rows[r_i + 1]
        prevent_row_split(rr)
        fill = "FFFFFF" if r_i % 2 == 0 else PALE
        for c_i, val in enumerate(row):
            cell = rr.cells[c_i]
            cell.text = ""
            shade_cell(cell, fill)
            set_cell_margins(cell, 36, 36, 70, 70)
            set_cell_borders(cell, "D5DAE0", "4")
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.space_before = Pt(0)
            run = para.add_run(str(val))
            set_run_font(run, "Calibri", 9, False, False, INK)
    return table
